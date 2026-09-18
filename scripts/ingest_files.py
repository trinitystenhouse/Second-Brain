#!/usr/bin/env python3
"""Turn whatever you dropped in an inbox into plain markdown Claude can atomise.

This script does the mechanical half of ingestion only: it finds files, pulls
readable text out of them, and writes staging files. It does NOT write your
notes — deciding what is one idea and what it links to is the part that needs
judgement, so Claude does that afterwards (say "process my inbox").

    python3 scripts/ingest_files.py .              # every inbox in the vault
    python3 scripts/ingest_files.py . --spoke Photography
    python3 scripts/ingest_files.py . --dry-run

Reads .md .txt .markdown .rst .csv .json out of the box, and .pdf .docx .html
.htm .epub .rtf where the relevant tool or library is available. It tells you
exactly what it could not read and what would fix it, then carries on.

Optional extras, all independent:
    pip install pypdf          # PDFs, no system tools needed
    pip install python-docx    # Word documents
    pip install beautifulsoup4 # better HTML extraction
System tools it will use if present: pdftotext (poppler), pandoc.

Exit codes: 0 all good, 1 some files could not be read, 2 usage error.
"""

import argparse
import datetime as dt
import html
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

PLAIN = {".md", ".markdown", ".txt", ".rst", ".csv", ".json", ".text"}
RICH = {".pdf", ".docx", ".html", ".htm", ".epub", ".rtf"}
SKIP_NAMES = {".DS_Store", ".gitkeep", "Thumbs.db"}
STAGED = "_staged"
PROCESSED = "_processed"


# ----------------------------------------------------------------- utilities

def slugify(text, maxlen=60):
    text = re.sub(r"[^\w\s-]", "", text.lower())
    text = re.sub(r"[\s_]+", "-", text).strip("-")
    return (text[:maxlen].rstrip("-")) or "untitled"


def have(cmd):
    return shutil.which(cmd) is not None


def strip_tags(markup):
    """Last-resort HTML to text, used when beautifulsoup4 is not installed."""
    markup = re.sub(r"(?is)<(script|style|nav|footer|header|aside)[^>]*>.*?</\1>", " ", markup)
    markup = re.sub(r"(?i)<br\s*/?>", "\n", markup)
    markup = re.sub(r"(?i)</(p|div|h[1-6]|li|tr)>", "\n\n", markup)
    markup = re.sub(r"(?s)<[^>]+>", " ", markup)
    markup = html.unescape(markup)
    markup = re.sub(r"[ \t]+", " ", markup)
    return re.sub(r"\n{3,}", "\n\n", markup).strip()


def html_to_text(markup):
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return strip_tags(markup), None
    soup = BeautifulSoup(markup, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else None
    text = re.sub(r"\n{3,}", "\n\n", soup.get_text("\n", strip=True))
    return text, title


# ------------------------------------------------------------- extractors
# Each returns (text, title_or_None) or raises Unreadable(reason, fix).

class Unreadable(Exception):
    def __init__(self, reason, fix=""):
        super().__init__(reason)
        self.reason, self.fix = reason, fix


def read_plain(path):
    try:
        return path.read_text(encoding="utf-8"), None
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="replace"), None


def read_pdf(path):
    if have("pdftotext"):
        out = subprocess.run(["pdftotext", "-layout", str(path), "-"],
                             capture_output=True, text=True, timeout=180)
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout, None
    try:
        from pypdf import PdfReader
    except ImportError:
        raise Unreadable("no PDF reader available", "pip install pypdf")
    try:
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:
        raise Unreadable(f"pypdf could not open it ({exc.__class__.__name__})", "")
    text = "\n\n".join(pages).strip()
    if not text:
        raise Unreadable("no extractable text — probably a scanned image",
                         "run it through OCR first, e.g. ocrmypdf")
    title = None
    try:
        meta = reader.metadata
        if meta and meta.title:
            title = str(meta.title).strip() or None
    except Exception:
        pass
    return text, title


def read_docx(path):
    try:
        import docx
    except ImportError:
        if have("pandoc"):
            out = subprocess.run(["pandoc", "-t", "markdown", str(path)],
                                 capture_output=True, text=True, timeout=180)
            if out.returncode == 0:
                return out.stdout, None
        raise Unreadable("no .docx reader available",
                         "pip install python-docx  (or install pandoc)")
    doc = docx.Document(str(path))
    parts = []
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            digits = re.findall(r"\d", style)
            level = min(int(digits[0]) if digits else 2, 3) + 1
            parts.append("#" * min(level, 3) + " " + txt)
        else:
            parts.append(txt)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("|", "\\|") for c in row.cells]
            parts.append("| " + " | ".join(cells) + " |")
    return "\n\n".join(parts), None


def read_html(path):
    return html_to_text(path.read_text(encoding="utf-8", errors="replace"))


def read_epub(path):
    try:
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist()
                     if n.lower().endswith((".xhtml", ".html", ".htm"))]
            if not names:
                raise Unreadable("no readable chapters inside the epub", "")
            chunks = []
            for name in sorted(names):
                raw = zf.read(name).decode("utf-8", errors="replace")
                text, _ = html_to_text(raw)
                if text.strip():
                    chunks.append(text)
    except zipfile.BadZipFile:
        raise Unreadable("not a valid epub archive", "")
    return "\n\n---\n\n".join(chunks), None


def read_rtf(path):
    if have("pandoc"):
        out = subprocess.run(["pandoc", "-f", "rtf", "-t", "markdown", str(path)],
                             capture_output=True, text=True, timeout=180)
        if out.returncode == 0:
            return out.stdout, None
    raise Unreadable("no .rtf reader available", "install pandoc")


EXTRACTORS = {
    ".pdf": read_pdf, ".docx": read_docx, ".html": read_html,
    ".htm": read_html, ".epub": read_epub, ".rtf": read_rtf,
}


# -------------------------------------------------------------------- core

def guess_title(text, fallback):
    for line in text.splitlines():
        line = line.strip().lstrip("#").strip()
        if 3 < len(line) < 120:
            return line
    return fallback


def find_inboxes(root, spoke=None):
    boxes = []
    if spoke:
        target = root / spoke / "00_Inbox"
        if not target.is_dir():
            print(f"error: {spoke}/00_Inbox does not exist", file=sys.stderr)
            sys.exit(2)
        return [target]
    if (root / "_inbox").is_dir():
        boxes.append(root / "_inbox")
    for child in sorted(root.iterdir()):
        if child.is_dir() and not child.name.startswith((".", "_")):
            box = child / "00_Inbox"
            if box.is_dir():
                boxes.append(box)
    return boxes


def candidate_files(box):
    for path in sorted(box.rglob("*")):
        if not path.is_file():
            continue
        if STAGED in path.parts or PROCESSED in path.parts:
            continue
        if path.name in SKIP_NAMES or path.name.startswith("."):
            continue
        if path.suffix.lower() in PLAIN or path.suffix.lower() in RICH:
            yield path


def stage(path, box, root, dry_run):
    ext = path.suffix.lower()
    if ext in PLAIN:
        text, title = read_plain(path)
    else:
        text, title = EXTRACTORS[ext](path)

    text = text.strip()
    if not text:
        raise Unreadable("file is empty once text is extracted", "")

    title = title or guess_title(text, path.stem.replace("-", " ").replace("_", " "))
    spoke = box.parent.name if box.name == "00_Inbox" else "unfiled"
    out_dir = box / STAGED
    out_name = f"{dt.date.today():%Y-%m-%d}-{slugify(title)}.md"
    out_path = out_dir / out_name

    words = len(text.split())
    header = (
        "---\n"
        f"ingested: {dt.date.today():%Y-%m-%d}\n"
        "type: staged\n"
        "status: awaiting-atomisation\n"
        f"original_file: {path.relative_to(root)}\n"
        f"original_format: {ext.lstrip('.')}\n"
        f"suggested_spoke: {spoke}\n"
        f"detected_title: {json.dumps(title)}\n"
        f"word_count: {words}\n"
        "url:\n"
        "---\n\n"
        "<!-- STAGED FOR ATOMISATION.\n"
        "     This is raw extracted text, not a note. Say \"process my inbox\"\n"
        "     to Claude and it will split this into proper source and permanent\n"
        "     notes, then move the original into _processed/.\n"
        "     If this came from a web page or a book, fill in the url/citation\n"
        "     above first so the source note can cite it honestly. -->\n\n"
    )

    if dry_run:
        return out_path, words, True

    out_dir.mkdir(parents=True, exist_ok=True)
    counter = 1
    while out_path.exists():
        out_path = out_dir / f"{out_name[:-3]}-{counter}.md"
        counter += 1
    out_path.write_text(header + text + "\n", encoding="utf-8")
    return out_path, words, False


def main():
    ap = argparse.ArgumentParser(
        description="Extract text from inbox files into staging notes for Claude to atomise.")
    ap.add_argument("root", nargs="?", default=".", help="vault root (default: current directory)")
    ap.add_argument("--spoke", help="only process this spoke's 00_Inbox")
    ap.add_argument("--dry-run", action="store_true", help="report what would happen, write nothing")
    args = ap.parse_args()

    root = Path(args.root).resolve()
    if not root.is_dir():
        print(f"error: {root} is not a directory", file=sys.stderr)
        return 2

    boxes = find_inboxes(root, args.spoke)
    if not boxes:
        print("No inbox folders found. Create a spoke first:")
        print("  bash scripts/new_spoke.sh MySpoke")
        return 0

    staged, failed, skipped = [], [], 0
    for box in boxes:
        for path in candidate_files(box):
            try:
                out_path, words, was_dry = stage(path, box, root, args.dry_run)
                staged.append((path.relative_to(root), out_path.relative_to(root), words))
            except Unreadable as exc:
                failed.append((path.relative_to(root), exc.reason, exc.fix))
            except Exception as exc:  # noqa: BLE001 - one bad file must not stop the run
                failed.append((path.relative_to(root), f"{exc.__class__.__name__}: {exc}", ""))

    verb = "would stage" if args.dry_run else "staged"
    print(f"{verb} {len(staged)} file(s) from {len(boxes)} inbox(es)\n")
    for src, dest, words in staged:
        print(f"  {src}")
        print(f"    -> {dest}  ({words:,} words)")

    if failed:
        print(f"\ncould not read {len(failed)} file(s):")
        for src, reason, fix in failed:
            print(f"  ! {src}")
            print(f"      {reason}")
            if fix:
                print(f"      fix: {fix}")

    if staged and not args.dry_run:
        print("\nNext: open Claude in this vault and say  process my inbox")
        print("It will atomise these into linked notes and clear the staging folder.")
    elif not staged and not failed:
        print("Nothing to do — every inbox is empty.")
        print("Drop files into _inbox/ or any spoke's 00_Inbox/ and run this again.")

    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
